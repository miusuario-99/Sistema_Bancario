from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from .models import Cuenta
from clientes.models import Cliente
from .forms import CuentaForm
import re
from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font
from django.http import HttpResponse
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from django.conf import settings
from datetime import datetime
import os
from auditoria.utils import registrar_auditoria
from usuarios.decorators import grupos_permitidos
from movimientos.models import Movimiento

# ==========================
# LISTAR CUENTAS
# ==========================

@grupos_permitidos("Administrador", "Supervisor", "Cajero")
def lista_cuentas(request):

    query = request.GET.get("q", "")

    # Obtener cuentas
    cuentas = Cuenta.objects.select_related(
        "cliente"
    ).order_by("-fecha_apertura")


    # BUSCADOR
    if query:

        cuentas = cuentas.filter(

            Q(numero_cuenta__icontains=query) |

            Q(cliente__nombres__icontains=query) |

            Q(cliente__apellidos__icontains=query)

        )

    # PAGINACIÓN
    paginator = Paginator(cuentas, 5)

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(page_number)

    # ESTADÍSTICAS

    total_cuentas = Cuenta.objects.count()

    cuentas_activas = Cuenta.objects.filter(
        estado="A"
    ).count()

    cuentas_inactivas = Cuenta.objects.filter(
        estado="I"
    ).count()

    saldo_total = Cuenta.objects.aggregate(
        total=Sum("saldo")
    )["total"] or 0

    return render(request, "cuentas/lista.html", {

    "page_obj": page_obj,

    "query": query,

    "total_cuentas": total_cuentas,

    "cuentas_activas": cuentas_activas,

    "saldo_total": saldo_total,

})

    # ==========================
    # NUEVA CUENTA
    # ==========================

@grupos_permitidos("Administrador")
def nueva_cuenta(request):

    # Buscar el mayor número de las cuentas con formato CTA-000001
    mayor = 0

    for cuenta in Cuenta.objects.all():

        m = re.match(r"CTA-(\d{6})$", cuenta.numero_cuenta)

        if m:

            numero = int(m.group(1))

            if numero > mayor:
                mayor = numero

    numero_generado = f"CTA-{mayor + 1:06d}"

    if request.method == "POST":

        form = CuentaForm(request.POST)

        if form.is_valid():

            cuenta = form.save(commit=False)
            cuenta.numero_cuenta = numero_generado
            cuenta.save()

            registrar_auditoria(
                request,
                "Cuentas",
                "Nueva cuenta",
                f"Cuenta {cuenta.numero_cuenta} creada para {cuenta.cliente}"
            )

            messages.success(
                request,
                "Cuenta creada correctamente."
            )

            return redirect("lista_cuentas")

    else:

        form = CuentaForm(
            initial={
                "numero_cuenta": numero_generado
            }
        )

    return render(
        request,
        "cuentas/formulario.html",
        {
            "form": form,
            "titulo": "Nueva Cuenta"
        }
    )

    # ==========================
    # EDITAR CUENTA
    # ==========================

@grupos_permitidos("Administrador")
def editar_cuenta(request, id):

    cuenta = get_object_or_404(
        Cuenta,
        id=id
    )


    if request.method == "POST":

        form = CuentaForm(
            request.POST,
            instance=cuenta
        )

        if form.is_valid():

            form.save()

            registrar_auditoria(
                request,
                "Cuentas",
                "Editar cuenta",
                f"Cuenta {cuenta.numero_cuenta}"
            )

            messages.success(
                request,
                "Cuenta actualizada correctamente."
            )

            return redirect(
                "lista_cuentas"
            )

    else:

        form = CuentaForm(
            instance=cuenta
        )

    return render(
        request,
        "cuentas/formulario.html",
        {
            "form": form,
            "titulo": "Editar Cuenta"
        }
    )

@grupos_permitidos("Administrador")
def eliminar_cuenta(request, id):

    if request.method != "POST":
        messages.error(
            request,
            "La eliminación requiere una solicitud POST."
        )
        return redirect("lista_cuentas")

    cuenta = get_object_or_404(Cuenta, id=id)

    # No eliminar cuentas que tengan historial bancario
    if Movimiento.objects.filter(cuenta=cuenta).exists():
        messages.error(
            request,
            "No se puede eliminar esta cuenta porque "
            "tiene movimientos registrados. "
            "Debe conservarse su historial bancario."
        )
        return redirect("lista_cuentas")

    numero_cuenta = cuenta.numero_cuenta
    cliente = cuenta.cliente

    registrar_auditoria(
        request,
        "Cuentas",
        "Eliminar cuenta",
        f"Cuenta {numero_cuenta} eliminada. Cliente: {cliente}"
    )

    cuenta.delete()

    messages.success(
        request,
        "Cuenta eliminada correctamente."
    )

    return redirect("lista_cuentas")
# ==========================
# EXPORTAR CUENTAS A EXCEL
# ==========================

@grupos_permitidos("Administrador", "Supervisor")
def exportar_cuentas_excel(request):

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from copy import copy

    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas"

    # ENCABEZADOS
    ws.append([
        "N.º",
        "Número de cuenta",
        "Cliente",
        "Tipo",
        "Saldo",
        "Fecha de apertura",
        "Estado",
    ])

    cuentas = (
        Cuenta.objects
        .select_related("cliente")
        .order_by("numero_cuenta")
    )

    for numero, cuenta in enumerate(cuentas, start=1):
        ws.append([
            numero,
            cuenta.numero_cuenta,
            f"{cuenta.cliente.nombres} {cuenta.cliente.apellidos}",
            cuenta.get_tipo_cuenta_display(),
            float(cuenta.saldo),
            cuenta.fecha_apertura.strftime("%d/%m/%Y")
                if cuenta.fecha_apertura else "",
            cuenta.get_estado_display(),
        ])

    ultima_fila = ws.max_row

    # FORMATO DEL ENCABEZADO
    for celda in ws[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="0B3A75"
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    ws.row_dimensions[1].height = 25

    # FORMATO DE LAS FILAS
    borde = Side(style="thin", color="D9E2F3")

    for fila in ws.iter_rows(
        min_row=2,
        max_row=ultima_fila
    ):
        for celda in fila:
            celda.border = Border(
                bottom=borde
            )
            celda.alignment = Alignment(
                vertical="center"
            )

        # Saldo en soles
        fila[4].number_format = '"S/." #,##0.00'

        # Centrar número, fecha y estado
        for indice in (0, 5, 6):
            fila[indice].alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # RESUMEN
    fila_resumen = ultima_fila + 2

    ws.cell(fila_resumen, 1, "Total de cuentas:")
    ws.cell(fila_resumen, 2, cuentas.count())

    ws.cell(fila_resumen + 1, 1, "Saldo total:")
    ws.cell(fila_resumen + 1, 5,
            float(cuentas.aggregate(
                total=Sum("saldo")
            )["total"] or 0))

    ws.cell(
        fila_resumen + 1, 5
    ).number_format = '"S/." #,##0.00'

    for fila in (fila_resumen, fila_resumen + 1):
        for celda in ws[fila]:
            celda.font = Font(bold=True)

    # ANCHOS DE COLUMNAS
    anchos = {
        "A": 8,
        "B": 22,
        "C": 38,
        "D": 22,
        "E": 18,
        "F": 22,
        "G": 16,
    }

    for columna, ancho in anchos.items():
        ws.column_dimensions[columna].width = ancho

    # FILTROS Y ENCABEZADO FIJO
    ws.auto_filter.ref = f"A1:G{ultima_fila}"
    ws.freeze_panes = "A2"

    # RESPUESTA EXCEL
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="cuentas.xlsx"'
    )

    wb.save(response)

    return response

# ==========================
# EXPORTAR CUENTAS PDF
# ==========================

@grupos_permitidos("Administrador", "Supervisor")
def exportar_cuentas_pdf(request):

    print(">>> EXPORTANDO CUENTAS PDF <<<")

    response = HttpResponse(content_type="application/pdf")

    response["Content-Disposition"] = (
        'attachment; filename="cuentas.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=letter
    )

    elementos = []

    estilos = getSampleStyleSheet()

    # ==========================
    # LOGO
    # ==========================

    ruta_logo = os.path.join(
        settings.BASE_DIR,
        "static",
        "img",
        "logo_banco.png"
    )

    if os.path.exists(ruta_logo):

        logo = Image(
            ruta_logo,
            width=80,
            height=80
        )

        elementos.append(logo)

    elementos.append(Spacer(1, 10))

    # ==========================
    # TITULO
    # ==========================

    elementos.append(
        Paragraph(
            "Banco Python",
            estilos["Title"]
        )
    )

    elementos.append(
        Paragraph(
            "Reporte General de Cuentas",
            estilos["Heading2"]
        )
    )

    fecha = datetime.now().strftime("%d/%m/%Y %H:%M")

    elementos.append(
        Paragraph(
            f"Fecha de emisión: {fecha}",
            estilos["Normal"]
        )
    )

    elementos.append(Spacer(1,20))

    # ==========================
    # TABLA
    # ==========================

    datos = [[
        "N°",
        "Cuenta",
        "Cliente",
        "Tipo",
        "Saldo",
        "Estado"
    ]]

    cuentas = list(
        Cuenta.objects.select_related("cliente")
    )

    cuentas.sort(
        key=lambda cuenta: int(
            cuenta.numero_cuenta.replace("CTA-", "")
        )
    )

    contador = 1

    for cuenta in cuentas:

        datos.append([
            contador,
            cuenta.numero_cuenta,
            f"{cuenta.cliente.nombres} {cuenta.cliente.apellidos}",
            cuenta.get_tipo_cuenta_display(),
            f"S/. {cuenta.saldo:,.2f}",
            cuenta.get_estado_display()
        ])

        contador += 1

    tabla = Table(datos, repeatRows=1)

    tabla.setStyle(TableStyle([

        ("BACKGROUND",(0,0),(-1,0),colors.darkblue),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("GRID",(0,0),(-1,-1),0.5,colors.grey),
        ("ALIGN",(0,0),(-1,-1),"CENTER"),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("BOTTOMPADDING",(0,0),(-1,0),10),

    ]))

    elementos.append(tabla)

    elementos.append(Spacer(1,20))

    elementos.append(

        Paragraph(

            f"Total de cuentas: {len(cuentas)}",

            estilos["Normal"]

        )

    )

    doc.build(elementos)

    return response
